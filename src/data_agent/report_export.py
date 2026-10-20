"""从分析报告 Markdown 生成 Word / PPT，供业务转发。"""

from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.util import Inches as PptInches
from pptx.util import Pt as PptPt


def _sections(markdown: str) -> list[tuple[str, str]]:
    text = (markdown or "").strip()
    if not text:
        return []
    parts = re.split(r"\n(?=## )", text)
    sections: list[tuple[str, str]] = []
    for part in parts:
        lines = part.strip().splitlines()
        if not lines:
            continue
        title = re.sub(r"^#+\s*", "", lines[0]).strip() or "报告"
        body = "\n".join(lines[1:]).strip()
        sections.append((title, body))
    return sections


def build_docx(title: str, markdown: str, chart_paths: list[Path] | None = None) -> bytes:
    document = Document()
    heading = document.add_heading(title or "数据分析报告", level=0)
    heading.runs[0].font.size = Pt(22)
    for section_title, body in _sections(markdown) or [("报告", markdown or "暂无内容")]:
        document.add_heading(section_title, level=1)
        for paragraph in body.split("\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
    for path in (chart_paths or [])[:8]:
        if path.suffix.lower() in {".png", ".jpg", ".jpeg"} and path.is_file():
            document.add_heading(path.stem, level=2)
            document.add_picture(str(path), width=Inches(6.2))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pptx(title: str, markdown: str, chart_paths: list[Path] | None = None) -> bytes:
    presentation = Presentation()
    blank = presentation.slide_layouts[6] if len(presentation.slide_layouts) > 6 else presentation.slide_layouts[0]
    title_layout = presentation.slide_layouts[0]
    title_slide = presentation.slides.add_slide(title_layout)
    title_slide.shapes.title.text = title or "数据分析报告"
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = "由 Data Analysis Agent 根据工具结果生成"

    sections = _sections(markdown)
    overview = next((body for name, body in sections if "结论" in name or "Executive" in name), "")
    if overview:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "结论速览"
        slide.placeholders[1].text = overview[:1800]

    for path in (chart_paths or [])[:3]:
        if not (path.suffix.lower() in {".png", ".jpg", ".jpeg"} and path.is_file()):
            continue
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), PptInches(0.4), PptInches(0.8), width=PptInches(9.2))
        box = slide.shapes.add_textbox(PptInches(0.4), PptInches(0.2), PptInches(9), PptInches(0.5))
        box.text_frame.text = path.stem
        box.text_frame.paragraphs[0].font.size = PptPt(16)

    caveats = next((body for name, body in sections if "注意" in name or "Caveat" in name), "")
    if caveats:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "局限与建议"
        slide.placeholders[1].text = caveats[:1800]

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
